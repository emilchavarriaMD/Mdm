"""Exportación de listados y reportes a Excel (OpenPyXL).

Solo se exportan las columnas necesarias para el propósito del reporte.
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="0F6E7C")
THIN = Side(style="thin", color="B8C7CC")


def export_report(path, report: dict, institucion="") -> str:
    """report: dict(title, headers, rows, notes, filters). Devuelve la ruta escrita."""
    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:31].replace("/", "-")
    ncols = max(len(report["headers"]), 1)

    ws.cell(row=1, column=1, value=report["title"]).font = Font(bold=True, size=14, color="0F6E7C")
    ws.cell(row=2, column=1, value=institucion or "").font = Font(size=10, color="5F6B70")
    ws.cell(row=3, column=1, value=report.get("filters", "")).font = Font(size=9, italic=True)
    ws.cell(row=4, column=1, value="Generado: " + datetime.now().strftime("%d/%m/%Y %H:%M")).font = \
        Font(size=9, italic=True, color="5F6B70")

    header_row = 6
    for c, h in enumerate(report["headers"], 1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.row_dimensions[header_row].height = 32

    for r, row in enumerate(report["rows"], header_row + 1):
        is_total = bool(row) and str(row[0]).upper() == "TOTAL"
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if is_total:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8F3F5")

    # anchos de columna
    for c in range(1, ncols + 1):
        values = [str(report["headers"][c - 1])] + [str(r[c - 1]) for r in report["rows"] if len(r) >= c]
        longest = max((len(v) for v in values), default=10)
        ws.column_dimensions[get_column_letter(c)].width = min(max(12, longest + 2), 48)

    note_row = header_row + len(report["rows"]) + 2
    for i, note in enumerate(report.get("notes", [])):
        ws.cell(row=note_row + i, column=1, value=note).font = Font(italic=True, size=9, color="5F6B70")

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if report["rows"]:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{header_row + len(report['rows'])}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(str(path))
    return str(path)
