"""Resumen de la brigada en Excel, a partir de los pacientes cargados desde los ATA."""
from collections import Counter

from services import ata as ata_svc
from services.coherencia import _anios
from services.validators import fmt_date

GRUPOS = [("Menores de 1 año", 0, 1), ("1 a 4 años", 1, 5), ("5 a 9 años", 5, 10), ("10 a 19 años", 10, 20),
          ("20 a 39 años", 20, 40), ("40 a 59 años", 40, 60), ("60 años o más", 60, 200)]


def grupo_edad(r):
    a = _anios(r.get("edad_anios"), r.get("edad_meses"), r.get("edad_dias"))
    if a is None:
        return "Sin edad"
    return next((n for n, lo, hi in GRUPOS if lo <= a < hi), "Sin edad")


def datos(db):
    rows = ata_svc.list_records(db)
    sexo = Counter({"M": 0, "H": 0})
    for r in rows:
        sexo[r["sexo"] or "?"] += 1
    edades = Counter(grupo_edad(r) for r in rows)
    estados = Counter(r["ai_estado"] for r in rows)
    dx = Counter()
    for r in rows:
        for k, c in (("dx1", "cie1"), ("dx2", "cie2"), ("dx3", "cie3")):
            d = (r.get(k) or "").strip()
            if d:
                dx[(d, (r.get(c) or "").strip())] += 1
    fechas = sorted({r["fecha"] for r in rows if r.get("fecha")})
    return {"rows": rows, "total": len(rows), "sexo": sexo, "edades": edades, "estados": estados, "dx": dx, "fechas": fechas}


def exportar(db, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    d = datos(db)
    wb = Workbook()
    head = PatternFill("solid", fgColor="1F4E78")

    def encabezado(ws, fila, cols):
        ws.append(cols)
        for i in range(1, len(cols) + 1):
            c = ws.cell(fila, i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head
            c.alignment = Alignment(horizontal="center", wrap_text=True)

    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Resumen de la brigada"])
    ws["A1"].font = Font(bold=True, size=14)
    if d["fechas"]:
        ws.append(["Fechas de atención", " a ".join(dict.fromkeys([fmt_date(d["fechas"][0]), fmt_date(d["fechas"][-1])]))])
    ws.append(["Pacientes atendidos", d["total"]])
    ws.append([])
    ws.append(["Por sexo"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["Mujeres", d["sexo"].get("M", 0)])
    ws.append(["Hombres", d["sexo"].get("H", 0)])
    if d["sexo"].get("?"):
        ws.append(["Sin dato", d["sexo"]["?"]])
    ws.append([])
    ws.append(["Por grupo de edad"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for n, _lo, _hi in GRUPOS + [("Sin edad", 0, 0)]:
        if d["edades"].get(n):
            ws.append([n, d["edades"][n]])
    ws.append([])
    ws.append(["Registro en ActivityInfo"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for e in ("Pendiente", "Llenado", "Registrado"):
        ws.append([e, d["estados"].get(e, 0)])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22

    w2 = wb.create_sheet("Diagnósticos")
    encabezado(w2, 1, ["Diagnóstico", "CIE", "Pacientes", "% de pacientes"])
    for (dx, cie), n in sorted(d["dx"].items(), key=lambda x: (-x[1], x[0][0])):
        w2.append([dx, cie, n, round(100 * n / d["total"], 1) if d["total"] else 0])
    for col, w in zip("ABCD", (44, 10, 11, 14)):
        w2.column_dimensions[col].width = w

    w3 = wb.create_sheet("Pacientes")
    cols = ["Fecha", "Paciente", "Sexo", "Edad", "Diagnóstico 1", "Diagnóstico 2", "Diagnóstico 3", "ActivityInfo"]
    encabezado(w3, 1, cols)
    for r in d["rows"]:
        w3.append([fmt_date(r["fecha"]) if r.get("fecha") else "", r["nombre"], r["sexo"], ata_svc.age_text(r),
                   r["dx1"], r["dx2"], r["dx3"], r["ai_estado"]])
    for i, w in enumerate((12, 32, 6, 10, 28, 28, 28, 14), 1):
        w3.column_dimensions[get_column_letter(i)].width = w
    w3.freeze_panes = "A2"
    wb.save(path)
    return d["total"]
