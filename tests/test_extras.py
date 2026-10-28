"""Papelera, alertas de incoherencias, resumen en Excel y bitácora del llenado."""
import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
_TMP = tempfile.mkdtemp(prefix="brigadas_extras_")
os.environ["BRIGADAS_DATA_DIR"] = _TMP

from database import Database  # noqa: E402
from services import ata as A  # noqa: E402
from services import coherencia, fill_log, resumen  # noqa: E402

HEADER = {"fecha": "2026-10-13", "establecimiento": "Brisas", "profesional": "Dra", "codigo": "1",
          "tipo_establecimiento": "", "departamento": "Cortés", "municipio": "Choloma"}


def fila(nombre, sexo, anios, dx1="", **kw):
    r = {k: "" for k in A.KEYS}
    r.update(nombre=nombre, sexo=sexo, edad_anios=str(anios), dx1=dx1, cie1="X1", **kw)
    return r


class T20_Extras(unittest.TestCase):
    def setUp(self):
        self.db = Database(Path(tempfile.mkdtemp(dir=_TMP)) / "t.db")
        self.sid = A.save_sheet(self.db, HEADER, [fila("Ana", "M", 34, "HTA"), fila("Luis", "H", 8, "Parasitosis"),
                                                  fila("Rosa", "M", 63, "HTA")])

    def test_papelera_recupera_y_no_pierde_datos(self):
        ids = [r["id"] for r in A.list_records(self.db)]
        A.set_registered(self.db, ids[0], True)
        self.assertEqual(A.delete_records(self.db, ids[:2]), 2)
        self.assertEqual(len(A.list_records(self.db)), 1)
        self.assertEqual(len(A.list_papelera(self.db)), 2)
        # editar la hoja no borra definitivamente lo que está en la papelera
        _sheet, rows = A.get_sheet(self.db, self.sid)
        self.assertEqual(len(rows), 1)
        A.save_sheet(self.db, HEADER, rows, sheet_id=self.sid)
        self.assertEqual(len(A.list_papelera(self.db)), 2)
        self.assertEqual(A.restore_records(self.db, [r["id"] for r in A.list_papelera(self.db)]), 2)
        rec = {r["nombre"]: r for r in A.list_records(self.db)}
        self.assertEqual(set(rec), {"Ana", "Luis", "Rosa"})
        self.assertEqual(rec["Ana"]["ai_estado"], "Pendiente" if rec["Ana"]["ai_estado"] != "Registrado" else "Registrado")

    def test_buscar_por_diagnostico(self):
        self.assertEqual({r["nombre"] for r in A.list_records(self.db, "hta")}, {"Ana", "Rosa"})

    def test_alertas(self):
        self.assertTrue(coherencia.alertas("H", 30, ["Vaginosis mixta"]))
        self.assertTrue(coherencia.alertas("M", 30, ["Hiperplasia prostática"]))
        self.assertTrue(coherencia.alertas("M", 40, ["Control niño sano"]))
        self.assertTrue(coherencia.alertas("M", 5, ["HTA"]))
        self.assertTrue(coherencia.alertas("M", 30, ["Dengue"], "36.5"))
        self.assertEqual(coherencia.alertas("M", 34, ["HTA", "Dermatitis", "Vaginosis mixta"], "36.6"), [])
        self.assertEqual(coherencia.alertas("H", 8, ["Parasitosis intestinal", "Faringitis aguda"]), [])
        self.assertEqual(coherencia.alertas("M", 24, ["PF (ACOs)", "Embarazo"]), [])
        self.assertEqual(coherencia.alertas("M", None, ["HTA"]), [])

    def test_resumen_excel(self):
        from openpyxl import load_workbook
        d = resumen.datos(self.db)
        self.assertEqual(d["total"], 3)
        self.assertEqual(d["sexo"]["M"], 2)
        self.assertEqual(d["dx"][("HTA", "X1")], 2)
        self.assertEqual(d["edades"]["5 a 9 años"], 1)
        out = Path(_TMP) / "res.xlsx"
        self.assertEqual(resumen.exportar(self.db, out), 3)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Resumen", "Diagnósticos", "Pacientes"])
        self.assertEqual(wb["Diagnósticos"]["A2"].value, "HTA")
        self.assertEqual(wb["Pacientes"].max_row, 4)

    def test_bitacora(self):
        p = fill_log.log_path()
        if p.exists():
            p.unlink()
        fill_log.registrar("Ana", {"ficha": ["Perfil: falta"], "aps": [], "ssr": [], "verif": ["«Plan» quedó vacío"]}, None, ["APS"])
        fill_log.registrar("Luis", None, RuntimeError("sin Chrome"), ["APS"])
        e = fill_log.leer()
        self.assertEqual([x["paciente"] for x in e], ["Luis", "Ana"])
        self.assertIn("sin Chrome", fill_log.texto(e[0]))
        self.assertIn("«Plan» quedó vacío", fill_log.texto(e[1]))


if __name__ == "__main__":
    unittest.main()


def _plantilla_ata_2023(path):
    """Hoja con la misma distribución de la plantilla ATA-2023 (encabezados de 3 filas, celdas combinadas, 20 filas)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["A5"], ws["D5"], ws["J5"], ws["M5"], ws["O5"] = "Establecimiento:", "Código:", "Tipo De Establecimiento:", "Departamento:", "Municipio:"
    ws["A6"], ws["C6"], ws["M6"] = "Profesional De Salud: MG", "Nombre: Dra Prueba", "Fecha (dd/mm/aa):"
    for rng in ("A5:B5", "D5:E5", "J5:K5", "P5:Q5", "A6:B6", "C6:L6", "A7:A9", "B7:B9", "C7:C9", "D7:D9", "E7:E9",
                "F7:H8", "I7:I9", "J7:L8", "M7:O7", "M8:M9", "N8:N9", "O8:O9", "P7:Q9"):
        ws.merge_cells(rng)
    ws["A7"], ws["B7"], ws["C7"], ws["D7"] = "N°", "NOMBRES Y APELLIDOS", "NÚMERO DE TELEFONO", "SEXO (M/F)"
    ws["E7"], ws["F7"], ws["I7"], ws["J7"] = "FECHA DE NACIMIENTO (dd/mm/aa)", "EDAD", "PACIENTE", "PROCEDENCIA"
    ws["M7"], ws["P7"] = "DIAGNÓSTICO / ACTIVIDAD", "Gestas, partos, cesareas, abortos, HV, HM"
    ws["M8"], ws["N8"], ws["O8"] = "1", "2", "3"
    ws["F9"], ws["G9"], ws["H9"], ws["J9"], ws["K9"], ws["L9"] = "AÑOS", "MESES", "DÍAS", "DEPARTAMENTO", "MUNICIPIO", "LOCALIDAD"
    for i in range(20):
        ws[f"A{10 + i}"] = i + 1
        ws.merge_cells(f"P{10 + i}:Q{10 + i}")
        ws[f"P{10 + i}"] = "G=    P=    C=    A=    HV=    HM=   "
    ws["C5"], ws["F5"], ws["L5"], ws["N5"], ws["P5"], ws["N6"] = "Brisas", "77", "Brigada", "Cortés", "Choloma", "15/10/26"
    ws["B10"], ws["C10"], ws["D10"], ws["E10"], ws["F10"] = "Ana Pérez", "9988-7766", "F", "12/03/90", 34
    ws["J10"], ws["K10"], ws["L10"], ws["M10"], ws["N10"], ws["P10"] = "Cortés", "Choloma", "Brisas 1", "HTA", "PF (ACOs)", "G=2 P=1 C=1 A=0 HV=3 HM=0"
    ws["B11"], ws["D11"], ws["G11"], ws["J11"], ws["K11"], ws["L11"], ws["M11"] = "Luis Mejía", "M", 6, '"', '"', '"', "Diarrea"
    wb.save(path)


class T21_PlantillaATA2023(unittest.TestCase):
    def test_lee_la_plantilla_con_encabezados_de_varias_filas(self):
        p = Path(tempfile.mkdtemp(dir=_TMP)) / "ata2023.xlsx"
        _plantilla_ata_2023(p)
        h, rows, warns = A.import_excel(p)
        self.assertEqual((h["establecimiento"], h["codigo"], h["municipio"], h["fecha"]), ("Brisas", "77", "Choloma", "15/10/2026"))
        self.assertEqual(h["profesional"], "Dra Prueba")
        self.assertEqual(len(rows), 2)                           # las 18 filas en blanco (con «G=  P=  …» vacío) no cuentan
        a, b = rows
        self.assertEqual((a["nombre"], a["telefono"], a["fecha_nacimiento"], a["edad_anios"]), ("Ana Pérez", "9988-7766", "", "34"))   # la fecha de nacimiento ya no se carga
        self.assertEqual((a["dep"], a["mun"], a["loc"], a["dx1"], a["dx2"]), ("Cortés", "Choloma", "Brisas 1", "HTA", "PF (ACOs)"))
        self.assertEqual(a["obst"], "G=2 P=1 C=1 A=0 HV=3 HM=0")
        self.assertEqual((a["sexo"], b["sexo"]), ("M", "H"))        # (M/F): F = mujer -> M, M = masculino -> H
        self.assertEqual((b["edad_meses"], b["loc"], b["obst"]), ("6", '"', ""))
        self.assertTrue(any("sexo" in w.lower() for w in warns))
        hh, clean, errores = A.prepare_sheet(h, rows)
        self.assertEqual(errores, [])
        self.assertEqual(clean[1]["loc"], "Brisas 1")               # la comilla se expande con la fila anterior
        db = Database(Path(tempfile.mkdtemp(dir=_TMP)) / "t2.db")
        sid = A.save_sheet(db, hh, clean)
        _s, guardadas = A.get_sheet(db, sid)
        self.assertEqual((guardadas[0]["telefono"], guardadas[0]["obst"]), ("9988-7766", "G=2 P=1 C=1 A=0 HV=3 HM=0"))
        self.assertEqual(guardadas[1]["loc"], "Brisas 1")

    def test_numeros_obstetricos_pasan_a_los_antecedentes_de_ssr(self):
        from services import activityinfo as ai
        d = ai.build_defaults(Database(Path(tempfile.mkdtemp(dir=_TMP)) / "t3.db"), {
            "nombre": "Ana", "sexo": "M", "fecha": "2026-10-15", "edad_anios": "30", "dx1": "PF",
            "obst": "G=2 P=1 C=1 A=0 HV=3 HM=0"})
        self.assertEqual([d[k] for k in ("ssr_gestaciones", "ssr_partos", "ssr_cesareas", "ssr_abortos", "ssr_hijos_vivos", "ssr_hijos_muertos")],
                         ["2", "1", "1", "0", "3", "0"])
        d0 = ai.build_defaults(Database(Path(tempfile.mkdtemp(dir=_TMP)) / "t4.db"), {
            "nombre": "Ana", "sexo": "M", "fecha": "2026-10-15", "edad_anios": "30", "dx1": "PF", "obst": "G=0 P=0 C=0 A=0 HV=0 HM=0"})
        self.assertEqual(d0["ssr_embarazo_anterior"], "No hubo embarazo")
        self.assertNotIn("ssr_gestaciones", ai.build_defaults(Database(Path(tempfile.mkdtemp(dir=_TMP)) / "t5.db"), {
            "nombre": "Ana", "sexo": "M", "fecha": "2026-10-15", "edad_anios": "30", "dx1": "PF"}))


class T22_ExportarEnPlantilla(unittest.TestCase):
    def test_exportar_y_volver_a_subir_conserva_todo(self):
        rows = [dict(fila("Ana Pérez", "M", 34, "HTA"), telefono="9988-7766", dx2="PF (ACOs)", dep="Cortés", mun="Choloma",
                     loc="Brisas 1", obst="G=2 P=1 C=1 A=0 HV=3 HM=0", fecha_nacimiento="1990-03-12"),
                dict(fila("Luis Mejía", "H", "", "Diarrea"), edad_meses="6", loc="Brisas 1")]
        out = Path(tempfile.mkdtemp(dir=_TMP)) / "salida.xlsx"
        A.export_excel(out, HEADER, rows)
        h, vuelta, warns = A.import_excel(out)
        self.assertEqual((h["establecimiento"], h["profesional"], h["municipio"], h["fecha"]), ("Brisas", "Dra", "Choloma", "13/10/2026"))
        self.assertEqual(len(vuelta), 2)
        self.assertEqual(vuelta[0]["obst"], "G=2 P=1 C=1 A=0 HV=3 HM=0")
        self.assertEqual((vuelta[0]["sexo"], vuelta[1]["sexo"]), ("M", "H"))
        self.assertEqual((vuelta[0]["telefono"], vuelta[0]["fecha_nacimiento"], vuelta[1]["edad_meses"]), ("9988-7766", "", "6"))
        self.assertEqual(warns, [])

    def test_fecha_de_nacimiento_se_ignora_y_la_edad_solo_se_calcula_si_falta(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["B1"], ws["C1"], ws["D1"], ws["E1"] = "Nombre completo", "Sexo", "Fecha de nacimiento", "Edad años", "Diagnóstico 1"
        ws["A2"], ws["B2"], ws["C2"], ws["D2"], ws["E2"] = "Con edad", "M", "12/03/1990", 34, "HTA"          # trae edad: se respeta
        ws["A3"], ws["B3"], ws["C3"], ws["D3"], ws["E3"] = "Sin edad", "H", "10/03/2018", None, "Diarrea"    # sin edad: se calcula
        ws["A4"], ws["B4"], ws["C4"], ws["D4"], ws["E4"] = "Sin nada", "H", None, None, "Gripe"
        ws["A5"], ws["B5"], ws["C5"], ws["D5"], ws["E5"] = "Fecha mala", "M", "99/99/9999", None, "Gripe"
        p = Path(tempfile.mkdtemp(dir=_TMP)) / "dob.xlsx"
        wb.save(p)
        h, rows, warns = A.import_excel(p)
        self.assertEqual([r["fecha_nacimiento"] for r in rows], ["", "", "", ""])
        self.assertEqual(rows[0]["edad_anios"], "34")
        hoy = date.today()
        esperada = hoy.year - 2018 - ((hoy.month, hoy.day) < (3, 10))
        self.assertEqual(rows[1]["edad_anios"], str(esperada))
        self.assertEqual((rows[2]["edad_anios"], rows[3]["edad_anios"]), ("", ""))
        self.assertTrue(any("1 fila" in w and "fecha de nacimiento" in w for w in warns), warns)

    def test_sexo_solo_se_convierte_si_la_hoja_usa_m_f(self):
        # en la hoja real el encabezado dice (M/F) pero se escribe H/M: no se cambia nada
        p = Path(tempfile.mkdtemp(dir=_TMP)) / "hm.xlsx"
        _plantilla_ata_2023(p)
        import openpyxl
        wb = openpyxl.load_workbook(p)
        ws = wb.active
        ws["D10"], ws["D11"] = "M", "H"
        wb.save(p)
        _h, rows, warns = A.import_excel(p)
        self.assertEqual([r["sexo"] for r in rows], ["M", "H"])
        self.assertEqual(warns, [])
        ws["D10"], ws["D11"] = "F", "M"       # M/F de verdad: F = mujer (M), M = masculino (H)
        wb.save(p)
        _h, rows, warns = A.import_excel(p)
        self.assertEqual([r["sexo"] for r in rows], ["M", "H"])
        self.assertTrue(warns)


class T23_NormalizarSexo(unittest.TestCase):
    def _n(self, *vals):
        rows = [{"sexo": v} for v in vals]
        A.normalizar_sexos(rows)
        return [r["sexo"] for r in rows]

    def test_palabras_y_letras(self):
        self.assertEqual(self._n("Masculino", "Femenino", "hombre", "MUJER", "masc.", "fem", "Varón"), ["H", "M", "H", "M", "H", "M", "H"])
        self.assertEqual(self._n("H", "F", "M"), ["H", "M", "M"])             # H/F: F es mujer; M es mujer
        self.assertEqual(self._n("H", "M", "h", "m"), ["H", "M", "H", "M"])
        self.assertEqual(self._n("F", "F", "M"), ["M", "M", "H"])             # M/F de verdad: M = masculino
        self.assertEqual(self._n("M", "M", "M", "F"), ["M", "M", "M", "M"])   # una «F» de error no voltea a las demás
        self.assertEqual(self._n("", "M"), ["", "M"])

    def test_se_corrige_al_guardar_y_al_subir(self):
        h = dict(HEADER, fecha="13/10/2026")
        rows = [dict(fila("Ana", "Femenino", 30), numero="1"), dict(fila("Luis", "Masculino", 40), numero="2")]
        _h, clean, errores = A.prepare_sheet(h, rows)
        self.assertEqual(errores, [])
        self.assertEqual([r["sexo"] for r in clean], ["M", "H"])
        p = Path(tempfile.mkdtemp(dir=_TMP)) / "sx.xlsx"
        _plantilla_ata_2023(p)
        import openpyxl
        wb = openpyxl.load_workbook(p)
        ws = wb.active
        ws["D10"], ws["D11"] = "Femenino", "Masculino"
        wb.save(p)
        _h, filas, warns = A.import_excel(p)
        self.assertEqual([r["sexo"] for r in filas], ["M", "H"])
        self.assertTrue(any("sexo" in w.lower() for w in warns))


class TestMedicamentos(unittest.TestCase):
    def test_catalogo_y_formato(self):
        from services import medicamentos as m
        self.assertGreaterEqual(len(m.MEDICAMENTOS), 40)
        self.assertEqual(len(set(m.MEDICAMENTOS)), len(m.MEDICAMENTOS))
        for t in m.MEDICAMENTOS:
            self.assertIn(":", t)
            self.assertRegex(t.split(":", 1)[1], r"\S")
        self.assertIn("Acetaminofén 500 mg tableta: 2 Tab VO c/6 horas por 5 días", m.MEDICAMENTOS)
        for insumo in ("Jeringa", "Guantes", "Gasa", "Algodón", "Esparadrapo", "Prueba de embarazo"):
            self.assertFalse(any(insumo.lower() in t.lower() for t in m.MEDICAMENTOS), insumo)
        claves = [m._clave(t) for t in m.MEDICAMENTOS]
        self.assertEqual(claves, sorted(claves))                                  # orden alfabético
        self.assertTrue(m.MEDICAMENTOS[0].startswith("Acetaminofén"))
        self.assertTrue(m.MEDICAMENTOS[-1].startswith("Trimetoprim"))
        self.assertEqual(m.agregar_al_plan("Plan.", []), "Plan.")
        t = m.agregar_al_plan("Plan.", ["A 1 Tab", "B 2 Tab"])
        self.assertEqual(t, "Plan.\nMedicamentos indicados: A 1 Tab; B 2 Tab.")
        self.assertEqual(m.agregar_al_plan(t, ["A 1 Tab", "B 2 Tab"]), t)     # no se repite


class TestChromeCerrado(unittest.TestCase):
    """Si la persona cierra la ventana de Chrome entre un registro y otro, el programa se reconecta solo."""

    class _Sw:
        def __init__(self, d):
            self.d = d

        def window(self, h):
            if h not in self.d.handles:
                raise RuntimeError("no such window: target window already closed")
            self.d.actual = h

    class _Drv:
        def __init__(self, handles, actual):
            self.handles, self.actual = list(handles), actual
            self.switch_to = TestChromeCerrado._Sw(self)
            self.service = type("S", (), {"stop": lambda s: None})()

        @property
        def window_handles(self):
            return list(self.handles)

        @property
        def current_url(self):
            if self.actual not in self.handles:
                raise RuntimeError("no such window: target window already closed from unknown error: web view not found")
            return "https://www.activityinfo.org/app"

    def test_sesion_sana(self):
        from services import activityinfo as ai
        self.assertTrue(ai._sesion_sana(self._Drv(["a", "b"], "b")))
        d = self._Drv(["a"], "b")                       # se cerró la pestaña a la que apuntaba, queda otra
        self.assertTrue(ai._sesion_sana(d))
        self.assertEqual(d.actual, "a")
        self.assertFalse(ai._sesion_sana(self._Drv([], "a")))     # no queda ninguna ventana

    def test_ventana_muerta(self):
        from services import activityinfo as ai
        for t in ("no such window: target window already closed", "web view not found", "invalid session id",
                  "disconnected: not connected to DevTools"):
            self.assertTrue(ai._ventana_muerta(Exception(t)), t)
        self.assertFalse(ai._ventana_muerta(Exception("No apareció el formulario")))

    def test_fill_form_reintenta_una_vez_con_chrome_nuevo(self):
        from services import activityinfo as ai
        llamadas = []
        guardados = (ai._get_driver, ai._fill_form_una_vez, ai._descartar_driver)
        try:
            ai._get_driver = lambda: llamadas.append("get") or "drv"
            ai._descartar_driver = lambda: llamadas.append("descartar")

            def una_vez(drv, data, espera):
                llamadas.append("llenar")
                if llamadas.count("llenar") == 1:
                    raise RuntimeError("no such window: target window already closed")
                return {"ficha": [], "aps": []}
            ai._fill_form_una_vez = una_vez
            self.assertEqual(ai.fill_form({}), {"ficha": [], "aps": []})
            self.assertEqual(llamadas, ["get", "llenar", "descartar", "get", "llenar"])
            llamadas.clear()
            ai._fill_form_una_vez = lambda *a: (_ for _ in ()).throw(RuntimeError("otro error"))   # otro error: no se reintenta
            with self.assertRaises(RuntimeError):
                ai.fill_form({})
            self.assertEqual(llamadas, ["get"])
        finally:
            ai._get_driver, ai._fill_form_una_vez, ai._descartar_driver = guardados


class TestATANuevo(unittest.TestCase):
    """ATA nuevo: estado civil en letras, profesión AM/O, una sola columna EDAD y localidad en el encabezado."""
    def _hoja(self):
        import openpyxl
        wb = openpyxl.load_workbook(Path(__file__).parent / "ata_nuevo_vacio.xlsx")
        ws = wb.active
        ws["C3"], ws["K3"], ws["D4"] = "Centro de Salud Sula", "Col. Las Brisas", "Fecha (dd/mm/aa): 23/10/26"
        filas = [("Ana Pérez", "9999-1111", "F", "UL", "AM", "34", "Candidiasis", "Parasitosis"),
                 ("Luis Gómez", "", "M", "s", "O", "8m", "Resfriado", ""),
                 ("Pedro", "", "M", "x", "Estudiante", "5a 3m", "Gastritis", "")]
        for i, f in enumerate(filas):
            for col, v in zip("BCDEFGHI", f):
                ws[f"{col}{8 + i}"] = v
        p = Path(tempfile.mkdtemp(dir=_TMP)) / "n.xlsx"
        wb.save(p)
        return p

    def test_importa_columnas_nuevas(self):
        h, rows, w = A.import_excel(self._hoja())
        self.assertEqual(len(rows), 3)
        self.assertEqual(h["departamento"], "Cortes")
        self.assertEqual(h["fecha"], "23/10/2026")
        self.assertEqual({r["loc"] for r in rows}, {"Col. Las Brisas"})
        self.assertEqual({r["dep"] for r in rows}, {"Cortes"})
        self.assertEqual((rows[0]["estado_civil"], rows[0]["profesion"], rows[0]["edad_anios"]), ("UL", "AM", "34"))
        self.assertEqual((rows[1]["estado_civil"], rows[1]["edad_meses"], rows[1]["edad_anios"]), ("S", "8", ""))
        self.assertEqual((rows[2]["edad_anios"], rows[2]["edad_meses"]), ("5", "3"))
        self.assertTrue(any("Estado civil no reconocido" in x for x in w))

    def test_equivalencias(self):
        self.assertEqual(A.estado_civil_ai("ul"), "Unión Libre")
        self.assertEqual(A.estado_civil_ai("C"), "Casado/a")
        self.assertEqual(A.estado_civil_ai("Soltera"), "Soltero/a")
        self.assertEqual(A.estado_civil_ai("zz"), "")
        self.assertEqual(A.ocupacion_ai("am"), "Trabajo doméstico NO remunerado")
        self.assertEqual(A.ocupacion_ai("O"), "Otro")
        self.assertEqual(A.ocupacion_ai("Estudiante"), "Estudiante")
        self.assertEqual(A.parse_edad("34"), ("34", "", ""))
        self.assertEqual(A.parse_edad("10 d"), ("", "", "10"))

    def test_guarda_y_llena_la_ficha(self):
        from services import activityinfo as ai
        db = Database(Path(tempfile.mkdtemp(dir=_TMP)) / "t.db")
        h, rows, _ = A.import_excel(self._hoja())
        h2, clean, errs = A.prepare_sheet(h, rows)
        self.assertEqual(errs, [])
        A.save_sheet(db, h2, clean)
        recs = {r["nombre"]: r for r in A.list_records(db)}
        d = ai.build_defaults(db, recs["Ana Pérez"])
        self.assertEqual((d["estado_civil"], d["ocupacion"]), ("Unión Libre", "Trabajo doméstico NO remunerado"))
        self.assertEqual(d["punto_especifico"], "Col. Las Brisas")
        d = ai.build_defaults(db, recs["Luis Gómez"])        # menor: «No aplica» manda
        self.assertEqual((d["estado_civil"], d["ocupacion"]), ("No aplica", "No aplica"))
        # exportar y volver a leer conserva estado civil y profesión
        out = Path(tempfile.mkdtemp(dir=_TMP)) / "e.xlsx"
        A.export_excel(out, h2, clean)
        _h, r2, _w = A.import_excel(out)
        self.assertEqual([(r["estado_civil"], r["profesion"]) for r in r2][:2], [("UL", "AM"), ("S", "O")])


class TestRevision(unittest.TestCase):
    def setUp(self):
        self.db = Database(Path(tempfile.mkdtemp(dir=_TMP)) / "t.db")

    def _hoja(self, nombre="Jose Lopes"):
        return A.save_sheet(self.db, HEADER, [fila(nombre, "H", 30, "HTA")])

    def test_corregir_nombre_conserva_registrado(self):
        sid = self._hoja()
        rid = A.list_records(self.db)[0]["id"]
        A.set_registered(self.db, rid, True)
        _s, rows = A.get_sheet(self.db, sid)
        nuevo = dict(fila("José López", "H", 30, "HTA"), _id=str(rows[0]["id"]))
        _h, clean, _e = A.prepare_sheet(HEADER, [nuevo])
        A.save_sheet(self.db, _h, clean, sid)
        self.assertEqual(A.list_records(self.db)[0]["ai_estado"], "Registrado")

    def test_busqueda_con_tildes(self):
        self._hoja("María GARCÍA")
        for q in ("garcía", "garcia", "MARIA", "maría garcía"):
            self.assertEqual(len(A.list_records(self.db, q)), 1, q)
        self.assertEqual(len(A.list_records(self.db, "peña")), 0)

    def test_edad_sin_dato_usa_fecha_nacimiento_y_cero_anios(self):
        from services import activityinfo as ai
        self._hoja()
        rec = dict(A.list_records(self.db)[0])
        rec.update(edad_anios="", fecha_nacimiento="2018-03-10", fecha="2026-10-23")
        self.assertEqual(ai.build_defaults(self.db, rec)["hea_persona"], {"sexo": "H", "edad": "8", "unidad": "años"})
        rec.update(edad_anios="0", fecha_nacimiento="")
        d = ai.build_defaults(self.db, rec)
        self.assertEqual(d["hea_persona"]["edad"], "0")
        self.assertEqual(d["estado_civil"], "No aplica")

    def test_edad_absurda_y_cervicalgia(self):
        _h, _c, errs = A.prepare_sheet(HEADER, [dict(fila("X", "H", 340, "HTA"))])
        self.assertTrue(any("no es válida" in e for e in errs))
        from services import coherencia
        self.assertFalse(coherencia._SOLO_MUJER.search("cervicalgia"))
        self.assertTrue(coherencia._SOLO_MUJER.search("cervicitis"))
