"""Pruebas automáticas de la capa de datos, reportes, PDF, Excel y respaldos.

Ejecutar:  python -m unittest tests.test_all -v
"""
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

_TMP = tempfile.mkdtemp(prefix="brigadas_test_")
os.environ["BRIGADAS_DATA_DIR"] = _TMP

import time
import backup  # noqa: E402
from database import Database  # noqa: E402
from reports import excel, pdf  # noqa: E402
from services import brigades as brigade_svc  # noqa: E402
from services import clinical, stats  # noqa: E402
from services import patients as P  # noqa: E402
from services import visits as V  # noqa: E402
from services.validators import parse_date, parse_number, ValidationError  # noqa: E402


def make_patient(db, nombres="María José", apellidos="Pérez López", **kw):
    raw = {"nombres": nombres, "apellidos": apellidos, "sexo": "Femenino", "edad_valor": "34",
           "edad_unidad": "años", "fecha_creacion": "10/03/2026", "departamento": "Yoro",
           "municipio": "El Progreso", "localidad": "Col. Centro", "telefono": "9988-7766",
           "estado_migratorio": "No migrante (nacional o residente)", **kw}
    data, errors = P.prepare_patient(raw)
    assert not errors, errors
    return P.create_patient(db, data)


def raw_visit(kind, fecha, **main):
    base = {"fecha_atencion": fecha}
    base.update(main)
    return {"main": base, "exam": {}, "dx": [], "meds": []}


class Base(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(dir=_TMP)
        self.db = Database(Path(self.dir) / "t.db")

    def tearDown(self):
        self.db.close()


class T01_Pacientes(Base):
    def test_01_crear_y_recuperar(self):
        pid = make_patient(self.db)
        p = P.get_patient(self.db, pid)
        self.assertEqual(p["codigo"], f"P{pid:06d}")
        self.assertEqual(p["nombres"], "María José")
        self.assertEqual(p["edad_valor"], 34)
        self.assertEqual(p["fecha_creacion"], "2026-03-10")
        self.assertEqual(p["telefono_norm"], "99887766")

    def test_02_editar_sin_perder_atenciones(self):
        pid = make_patient(self.db)
        data, _, _ = V.prepare_visit("aps", raw_visit("aps", "10/03/2026", personal="Dr. A"))
        vid = V.save_visit(self.db, "aps", pid, data)
        raw = {"nombres": "María José", "apellidos": "Pérez Gómez", "sexo": "Femenino",
               "edad_valor": "35", "edad_unidad": "años", "fecha_creacion": "10/03/2026"}
        new, errors = P.prepare_patient(raw)
        self.assertFalse(errors)
        changes = P.update_patient(self.db, pid, new)
        self.assertTrue(any("Apellidos" in c for c in changes))
        self.assertEqual(P.get_patient(self.db, pid)["apellidos"], "Pérez Gómez")
        self.assertIsNotNone(V.get_visit(self.db, "aps", vid))
        self.assertEqual(P.patient_counts(self.db, pid), (1, 0))

    def test_03_duplicados(self):
        pid = make_patient(self.db)
        dup, _ = P.prepare_patient({"nombres": "Maria Jose", "apellidos": "Perez Lopez",
                                    "sexo": "Femenino", "fecha_creacion": "11/03/2026"})
        found = P.find_duplicates(self.db, dup)
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0][0]["id"], pid)
        self.assertIn("mismo nombre", found[0][1])
        otro, _ = P.prepare_patient({"nombres": "Carlos", "apellidos": "Ramos", "sexo": "Masculino",
                                     "fecha_creacion": "11/03/2026", "telefono": "9988 7766"})
        self.assertIn("mismo teléfono", P.find_duplicates(self.db, otro)[0][1])
        distinto, _ = P.prepare_patient({"nombres": "Luis", "apellidos": "Zelaya", "sexo": "Masculino",
                                         "fecha_creacion": "11/03/2026"})
        self.assertEqual(P.find_duplicates(self.db, distinto), [])

    def test_04_edad_desde_fecha_nacimiento(self):
        data, errors = P.prepare_patient({"nombres": "Ana", "apellidos": "Díaz", "sexo": "Femenino",
                                          "fecha_nacimiento": "15/06/2020", "fecha_creacion": "10/03/2026"})
        self.assertFalse(errors)
        self.assertEqual((data["edad_valor"], data["edad_unidad"]), (5, "años"))
        bebe, _ = P.prepare_patient({"nombres": "Bebé", "apellidos": "Díaz", "sexo": "Masculino",
                                     "fecha_nacimiento": "20/12/2025", "fecha_creacion": "10/03/2026"})
        self.assertEqual((bebe["edad_valor"], bebe["edad_unidad"]), (2, "meses"))

    def test_05_busqueda(self):
        a = make_patient(self.db)
        make_patient(self.db, "Carlos", "Ramos", historia_clinica="HC-77", telefono="3322-1100")
        self.assertEqual(len(P.search_patients(self.db, "perez")), 1)
        self.assertEqual(len(P.search_patients(self.db, "HC-77")), 1)
        self.assertEqual(len(P.search_patients(self.db, "33221100")), 1)
        self.assertEqual(len(P.search_patients(self.db, P.get_patient(self.db, a)["codigo"])), 1)
        self.assertEqual(len(P.search_patients(self.db, "jose perez")), 1)
        self.assertEqual(len(P.search_patients(self.db)), 2)
        self.assertEqual(len(P.search_patients(self.db, tipo_ficha="Solo Ficha Integral (sin atenciones)")), 2)


class T02_Atenciones(Base):
    def _aps(self, pid, fecha, finalizar=False, **extra):
        raw = raw_visit("aps", fecha, personal="Dra. Ruiz", **extra)
        raw["dx"] = [{"tipo": "principal", "descripcion": "Faringitis aguda", "cie11": "CA02.0"}]
        data, errors, _ = V.prepare_visit("aps", raw, finalizing=finalizar)
        self.assertFalse(errors, errors)
        return V.save_visit(self.db, "aps", pid, data)

    def _ssr(self, pid, fecha, finalizar=False, **extra):
        raw = raw_visit("ssr", fecha, personal="Dra. Ruiz", motivo="Control", categoria="Planificación familiar", **extra)
        raw["dx"] = [{"tipo": "principal", "descripcion": "Control de planificación familiar"}]
        data, errors, _ = V.prepare_visit("ssr", raw, finalizing=finalizar)
        self.assertFalse(errors, errors)
        return V.save_visit(self.db, "ssr", pid, data)

    def test_06_aps_y_ssr_fechas_distintas(self):
        pid = make_patient(self.db)
        a = self._aps(pid, "10/03/2026")
        s = self._ssr(pid, "20/03/2026")
        self.assertEqual(V.suggest_same_day(self.db, "ssr", pid, "2026-03-20"), "No")
        self.assertEqual(V.get_visit(self.db, "aps", a)["main"]["fecha_atencion"], "2026-03-10")
        self.assertEqual(V.get_visit(self.db, "ssr", s)["main"]["fecha_atencion"], "2026-03-20")
        self.assertEqual(P.patient_counts(self.db, pid), (1, 1))

    def test_07_aps_y_ssr_mismo_dia(self):
        pid = make_patient(self.db)
        self.assertEqual(V.suggest_same_day(self.db, "ssr", pid, "2026-03-10"), "No se registró atención APS")
        self.assertEqual(V.suggest_same_day(self.db, "aps", pid, "2026-03-10"), "Sí")
        self._aps(pid, "10/03/2026")
        self.assertEqual(V.suggest_same_day(self.db, "ssr", pid, "2026-03-10"), "Sí")
        s = self._ssr(pid, "10/03/2026", mismo_dia_aps="Sí")
        self.assertEqual(V.get_visit(self.db, "ssr", s)["main"]["mismo_dia_aps"], "Sí")
        self.assertEqual(V.suggest_same_day(self.db, "aps", pid, "2026-04-01"), "No")

    def test_08_varias_atenciones_mismo_paciente_sin_duplicar(self):
        pid = make_patient(self.db)
        for f in ("10/03/2026", "17/03/2026", "24/03/2026"):
            self._aps(pid, f)
        self._ssr(pid, "17/03/2026")
        self._ssr(pid, "24/03/2026")
        self.assertEqual(P.patient_counts(self.db, pid), (3, 2))
        self.assertEqual(P.count_patients(self.db), 1)
        rows = V.list_all_visits(self.db, patient_id=pid)
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[0]["fecha_atencion"], "2026-03-24")

    def test_09_campos_clinicos_vacios_no_se_inventan(self):
        pid = make_patient(self.db)
        data, errors, warnings = V.prepare_visit("aps", raw_visit("aps", "10/03/2026"))
        self.assertFalse(errors)
        vid = V.save_visit(self.db, "aps", pid, data)
        v = V.get_visit(self.db, "aps", vid)
        for k in ("pas", "pad", "fc", "fr", "temperatura", "peso"):
            self.assertIsNone(v["main"][k], k)
        self.assertFalse(v["main"]["hea_texto"])
        self.assertEqual(v["dx"], [])
        self.assertEqual(v["meds"], [])
        self.assertEqual(v["main"]["estado"], "borrador")
        self.assertEqual(v["main"]["consejeria_estado"], "Pendiente")
        # examen: nada se marca como normal
        raw = raw_visit("aps", "10/03/2026")
        raw["exam"] = {"abdomen": {"estado": V.EXAM_PENDIENTE, "hallazgos": ""}}
        data, _, _ = V.prepare_visit("aps", raw)
        vid2 = V.save_visit(self.db, "aps", pid, data)
        exam = V.get_visit(self.db, "aps", vid2)["exam"]
        self.assertNotIn(V.EXAM_NORMAL, [e["estado"] for e in exam.values()])

    def test_10_varios_medicamentos(self):
        pid = make_patient(self.db)
        raw = raw_visit("aps", "10/03/2026", personal="Dr. X")
        raw["meds"] = [
            {"generico": "Acetaminofén", "presentacion": "Tabletas 500 mg", "dosis": "500 mg", "via": "Oral",
             "frecuencia": "Cada 8 horas", "horario": "", "duracion": "3 días", "cant_prescrita": "9",
             "cant_entregada": "9", "unidad": "tabletas", "indicaciones": "Con alimentos"},
            {"generico": "Amoxicilina", "presentacion": "Cápsulas 500 mg", "dosis": "500 mg", "via": "Oral",
             "frecuencia": "Cada 8 horas", "duracion": "7 días", "cant_prescrita": "21", "cant_entregada": "10",
             "unidad": "cápsulas"},
            {"generico": "Suero oral", "cant_prescrita": "3", "cant_entregada": "", "unidad": "sobres"},
            {"generico": "", "dosis": ""},  # fila vacía: se ignora
        ]
        data, errors, warnings = V.prepare_visit("aps", raw)
        self.assertFalse(errors)
        vid = V.save_visit(self.db, "aps", pid, data)
        meds = V.get_visit(self.db, "aps", vid)["meds"]
        self.assertEqual([m["generico"] for m in meds], ["Acetaminofén", "Amoxicilina", "Suero oral"])
        self.assertEqual(meds[1]["cant_prescrita"], 21.0)
        self.assertEqual(meds[1]["cant_entregada"], 10.0)
        self.assertIsNone(meds[2]["cant_entregada"], "entregada vacía debe quedar sin registrar")
        # entregada > prescrita: advertencia, no error
        raw["meds"][0]["cant_entregada"] = "12"
        _, errors, warnings = V.prepare_visit("aps", raw)
        self.assertFalse(errors)
        self.assertTrue(any("mayor que la prescrita" in w for w in warnings))
        # medicamento sin nombre pero con datos: error
        raw["meds"].append({"generico": "", "dosis": "10 mg"})
        _, errors, _ = V.prepare_visit("aps", raw)
        self.assertTrue(any("nombre genérico" in e for e in errors))

    def test_16_borrador_y_finalizacion(self):
        pid = make_patient(self.db)
        raw = raw_visit("aps", "10/03/2026")
        _, errors, _ = V.prepare_visit("aps", raw, finalizing=True)
        self.assertTrue(any("Personal médico" in e for e in errors))
        self.assertTrue(any("Diagnóstico principal" in e for e in errors))
        # HEA de plantilla sin revisar bloquea la finalización
        raw = raw_visit("aps", "10/03/2026", personal="Dr. X", hea_origen="plantilla", hea_revisada=0)
        raw["dx"] = [{"tipo": "principal", "descripcion": "Gastritis"}]
        _, errors, _ = V.prepare_visit("aps", raw, finalizing=True)
        self.assertTrue(any("revisada" in e for e in errors))
        raw["main"]["hea_revisada"] = 1
        data, errors, warnings = V.prepare_visit("aps", raw, finalizing=True)
        self.assertFalse(errors)
        self.assertTrue(any("examen físico" in w for w in warnings))
        vid = V.save_visit(self.db, "aps", pid, data)
        v = V.get_visit(self.db, "aps", vid)
        self.assertEqual(v["main"]["estado"], "finalizada")
        self.assertTrue(v["main"]["finalizada_at"])
        with self.assertRaises(ValueError):
            V.delete_draft(self.db, "aps", vid)  # las finalizadas no se eliminan

    def test_17_historial_de_modificaciones(self):
        pid = make_patient(self.db)
        data, _, _ = V.prepare_visit("aps", raw_visit("aps", "10/03/2026", personal="Dr. A"))
        vid = V.save_visit(self.db, "aps", pid, data)
        raw = raw_visit("aps", "10/03/2026", personal="Dr. B", pas="120")
        raw["dx"] = [{"tipo": "principal", "descripcion": "Gastritis"}]
        data2, _, _ = V.prepare_visit("aps", raw)
        V.save_visit(self.db, "aps", pid, data2, visit_id=vid)
        log = V.visit_audit(self.db, "aps", vid)
        self.assertEqual([r["accion"] for r in log], ["creación", "modificación"])
        self.assertIn("Dr. A", log[1]["detalle"])
        self.assertIn("diagnósticos modificado", log[1]["detalle"])
        # editar no altera otras atenciones
        otra = V.save_visit(self.db, "aps", pid, V.prepare_visit("aps", raw_visit("aps", "11/03/2026"))[0])
        antes = V.get_visit(self.db, "aps", otra)["main"]["updated_at"]
        V.save_visit(self.db, "aps", pid, data, visit_id=vid)
        self.assertEqual(V.get_visit(self.db, "aps", otra)["main"]["updated_at"], antes)

    def test_18_ssr_completa(self):
        pid = make_patient(self.db)
        raw = raw_visit("ssr", "10/03/2026", personal="Dra. R", motivo="Desea iniciar método",
                        categoria="Inicio de método anticonceptivo", fum="01/03/2026",
                        metodo_entregado="Inyectable trimestral", metodo_cantidad="1", metodo_unidad="dosis",
                        fecha_administracion="10/03/2026", fecha_proxima_dosis="02/06/2026")
        raw["dx"] = [{"tipo": "principal", "descripcion": "Planificación familiar"}]
        raw["tests"] = [{"prueba": "Prueba de embarazo", "estado": "Realizada", "resultado": "Negativa"},
                        {"prueba": "Examen general de orina", "estado": "No realizada", "motivo": "Sin insumos"}]
        raw["interventions"] = {V.SSR_INTERVENCIONES[0]: {"estado": "Realizada"},
                                V.SSR_INTERVENCIONES[2]: {"estado": "No realizada", "nota": "Sin existencias"}}
        data, errors, _ = V.prepare_visit("ssr", raw, finalizing=True)
        self.assertFalse(errors, errors)
        vid = V.save_visit(self.db, "ssr", pid, data)
        v = V.get_visit(self.db, "ssr", vid)
        self.assertEqual(v["main"]["fum"], "2026-03-01")
        self.assertEqual(len(v["tests"]), 2)
        self.assertEqual(v["interventions"][V.SSR_INTERVENCIONES[0]]["estado"], "Realizada")
        self.assertEqual(v["interventions"][V.SSR_INTERVENCIONES[1]]["estado"], "Pendiente")
        self.assertEqual(v["interventions"][V.SSR_INTERVENCIONES[2]]["estado"], "No realizada")


class T03_Validacion(Base):
    def test_15_fechas_numeros_obligatorios(self):
        self.assertEqual(parse_date("05/01/2026"), "2026-01-05")
        for bad in ("31/02/2026", "2026/13/01", "abc", "1/1/1800"):
            with self.assertRaises(ValidationError):
                parse_date(bad, "Fecha")
        with self.assertRaises(ValidationError):
            parse_date("01/01/2999", "Fecha", allow_future=False)
        self.assertEqual(parse_number("36,6"), 36.6)
        self.assertIsNone(parse_number("  "))
        with self.assertRaises(ValidationError):
            parse_number("doce")
        # paciente: obligatorios
        _, errors = P.prepare_patient({"nombres": "", "apellidos": "", "sexo": "", "fecha_creacion": ""})
        text = " ".join(errors)
        for needle in ("Nombres", "Apellidos", "Sexo", "Fecha de creación"):
            self.assertIn(needle, text)
        _, errors = P.prepare_patient({"nombres": "A", "apellidos": "B", "sexo": "Femenino",
                                       "fecha_creacion": "10/03/2026", "fecha_nacimiento": "31/12/2026"})
        self.assertTrue(errors)  # nacimiento posterior a la creación / futuro
        _, errors = P.prepare_patient({"nombres": "A", "apellidos": "B", "sexo": "Femenino",
                                       "fecha_creacion": "10/03/2026", "edad_valor": "abc"})
        self.assertTrue(any("Edad" in e for e in errors))
        _, errors = P.prepare_patient({"nombres": "A", "apellidos": "B", "sexo": "Femenino",
                                       "fecha_creacion": "10/03/2026", "edad_valor": "200"})
        self.assertTrue(any("Edad" in e for e in errors))
        _, errors = P.prepare_patient({"nombres": "A", "apellidos": "B", "sexo": "Femenino",
                                       "fecha_creacion": "10/03/2026", "telefono": "12"})
        self.assertTrue(any("Teléfono" in e for e in errors))
        # atención: fecha obligatoria, futura inválida, signos vitales
        _, errors, _ = V.prepare_visit("aps", raw_visit("aps", ""))
        self.assertTrue(any("Fecha de atención" in e for e in errors))
        _, errors, _ = V.prepare_visit("aps", raw_visit("aps", "01/01/2999"))
        self.assertTrue(any("futura" in e for e in errors))
        _, errors, _ = V.prepare_visit("aps", raw_visit("aps", "10/03/2026", pas="abc"))
        self.assertTrue(any("número" in e for e in errors))
        _, errors, _ = V.prepare_visit("aps", raw_visit("aps", "10/03/2026", temperatura="-5"))
        self.assertTrue(errors)
        _, errors, _ = V.prepare_visit("aps", raw_visit("aps", "10/03/2026", fecha_control="xx"))
        self.assertTrue(errors)
        # valores extremos pero posibles: advertencia, NO error
        data, errors, warnings = V.prepare_visit("aps", raw_visit("aps", "10/03/2026", pas="270", pad="90",
                                                                    temperatura="43.5"))
        self.assertFalse(errors)
        self.assertEqual(len([w for w in warnings if "rango habitual" in w]), 2)
        self.assertEqual(data["main"]["pas"], 270.0)
        # PAD > PAS: advertencia
        _, errors, warnings = V.prepare_visit("aps", raw_visit("aps", "10/03/2026", pas="80", pad="120"))
        self.assertTrue(any("diastólica es mayor" in w for w in warnings))


class T04_Clinico(Base):
    def test_hea_no_inventa_datos(self):
        p = {"sexo": "Masculino", "edad_valor": 40, "edad_unidad": "años", "fecha_nacimiento": None}
        text, pending = clinical.build_hea_aps(p, {}, "Faringitis aguda")
        self.assertIn("BORRADOR PRELIMINAR", text)
        self.assertIn("PENDIENTES", text)
        import re
        for inventado in ("fiebre", "tos", "odinofagia", "dolor de garganta", "días de evolución"):
            self.assertIsNone(re.search(r"\b" + inventado + r"\b", text.lower()), inventado)
        self.assertIn("Motivo de consulta", pending)
        text2, pending2 = clinical.build_hea_aps(p, {"hea_motivo": "Dolor de garganta", "hea_inicio": "2 días"},
                                                 "Faringitis aguda")
        self.assertIn("Dolor de garganta", text2)
        self.assertIn("2 días", text2)
        self.assertNotIn("Motivo de consulta", pending2)
        self.assertIn("Síntomas principales", pending2)

    def test_hea_ssr_no_inventa(self):
        p = {"sexo": "Femenino", "edad_valor": 25, "edad_unidad": "años"}
        text, pending = clinical.build_hea_ssr(p, {}, "Flujo vaginal", "Flujo vaginal anormal", "Vaginosis")
        self.assertIn("pendientes de documentar", text)
        self.assertIn("Características del flujo vaginal", pending)
        self.assertIn("Fecha de última menstruación (FUM)", pending)
        self.assertNotIn("prurito intenso", text.lower())

    def test_sugerencias(self):
        title, qs = clinical.suggest_questions("aps", "Infección de vías urinarias")
        self.assertIn("urinaria", title.lower())
        self.assertTrue(any("Disuria" in q for q in qs))
        signs = clinical.suggest_alarm_signs("aps", "Dengue sin signos de alarma")
        self.assertTrue(any("sangrado" in s.lower() for s in signs))
        title, qs = clinical.suggest_questions("aps", "Diagnóstico desconocido xyz")
        self.assertIn("generales", title.lower())
        # SSR: 'its' no debe coincidir dentro de otras palabras
        title, _ = clinical.suggest_questions("ssr", "Habits sociales")
        self.assertIn("generales", title.lower())


class T05_Reportes(Base):
    def _populate(self):
        b1 = brigade_svc.save_brigade(self.db, brigade_svc.prepare_brigade(
            {"codigo": "B-001", "nombre": "Brigada Yoro", "fecha": "10/03/2026", "departamento": "Yoro"})[0])
        b2 = brigade_svc.save_brigade(self.db, brigade_svc.prepare_brigade(
            {"codigo": "B-002", "nombre": "Brigada Cortés", "fecha": "20/03/2026"})[0])
        ids = []
        for n, (nom, ape) in enumerate([("Ana", "López"), ("Beto", "Mejía"), ("Carla", "Núñez")]):
            ids.append(make_patient(self.db, nom, ape, brigade_id=b1, telefono=f"9900-00{n}0"))
        def aps(pid, fecha, dx, b, meds=None):
            raw = raw_visit("aps", fecha, personal="Dr. A", brigade_id=str(b))
            raw["dx"] = [{"tipo": "principal", "descripcion": dx}]
            raw["meds"] = meds or []
            d, e, _ = V.prepare_visit("aps", raw, finalizing=False)
            assert not e, e
            return V.save_visit(self.db, "aps", pid, d)
        aps(ids[0], "10/03/2026", "Faringitis", b1, [{"generico": "Acetaminofén", "cant_prescrita": "6", "cant_entregada": "6", "unidad": "tabletas"}])
        aps(ids[0], "11/03/2026", "Faringitis", b1, [{"generico": "acetaminofén", "cant_prescrita": "6", "unidad": "tabletas"}])
        aps(ids[1], "10/03/2026", "Gastritis", b1)
        raw = raw_visit("ssr", "10/03/2026", personal="Dra. R", brigade_id=str(b1), motivo="x", categoria="Planificación familiar")
        raw["dx"] = [{"tipo": "principal", "descripcion": "Planificación familiar"}]
        V.save_visit(self.db, "ssr", ids[0], V.prepare_visit("ssr", raw)[0])
        aps(ids[2], "20/03/2026", "faringitis", b2)
        return b1, b2, ids

    def test_reportes_unicos_vs_consultas(self):
        b1, b2, ids = self._populate()
        r = stats.run_report(self.db, "Pacientes atendidos por brigada")
        by = {row[0]: row for row in r["rows"]}
        # B-001: 2 pacientes únicos (Ana, Beto); 3 APS + 1 SSR = 4 consultas
        self.assertEqual(by["B-001"][4], 2)
        self.assertEqual(by["B-001"][5], 3)
        self.assertEqual(by["B-001"][6], 1)
        self.assertEqual(by["B-001"][7], 4)
        self.assertEqual(by["B-002"][4], 1)
        r = stats.run_report(self.db, "Pacientes atendidos por día")
        total = r["rows"][-1]
        self.assertEqual(total[0], "TOTAL")
        self.assertEqual(total[1], 3)          # 3 pacientes únicos
        self.assertEqual(total[4], 5)          # 5 consultas
        r = stats.run_report(self.db, "Diagnósticos registrados")
        dx = {row[0].lower(): row for row in r["rows"]}
        self.assertEqual(dx["faringitis"][4], 3)   # agrupa 'Faringitis'/'faringitis'
        self.assertEqual(dx["faringitis"][5], 2)   # pacientes únicos
        r = stats.run_report(self.db, "Medicamentos prescritos y entregados")
        med = r["rows"][0]
        self.assertEqual(med[3], 2)            # 2 prescripciones
        self.assertEqual(med[4], "12")         # total prescrito
        self.assertEqual(med[5], 1)            # 1 con entrega documentada
        self.assertEqual(med[6], "6")          # solo lo documentado
        r = stats.run_report(self.db, "Atenciones por personal médico")
        self.assertEqual({row[0]: row[3] for row in r["rows"]}, {"Dr. A": 4, "Dra. R": 1})
        r = stats.run_report(self.db, "Pacientes atendidos por día", brigade_id=b2)
        self.assertEqual(r["rows"][-1][4], 1)
        r = stats.run_report(self.db, "Listado administrativo de pacientes")
        self.assertNotIn("Teléfono", r["headers"])
        self.assertEqual(len(r["rows"]), 3)
        r = stats.run_report(self.db, "Listado de atenciones")
        self.assertEqual(len(r["rows"]), 5)

    def test_panel(self):
        b1, b2, ids = self._populate()
        s = stats.dashboard_stats(self.db, b1)
        self.assertEqual(s["pacientes"], 3)
        self.assertEqual(s["brigada_atendidos"], 2)
        self.assertEqual(s["brigada_consultas"], 4)
        self.assertEqual(s["brigada_registrados"], 3)
        rows = brigade_svc.brigade_patients(self.db, b1)
        self.assertEqual(len(rows), 3)
        ana = [r for r in rows if r["nombres"] == "Ana"][0]
        self.assertEqual((ana["aps_n"], ana["ssr_n"]), (2, 1))
        self.assertIn("Faringitis", ana["diagnosticos"])
        with self.assertRaises(ValueError):
            brigade_svc.delete_brigade(self.db, b1)

    def test_11_pdf_tres_fichas(self):
        b1, b2, ids = self._populate()
        out = Path(self.dir)
        # APS completa con examen, vitales y medicamentos
        raw = raw_visit("aps", "10/03/2026", personal="Dr. A", brigade_id=str(b1), pas="120", pad="80",
                        temperatura="37.2", hea_motivo="Dolor de garganta", hea_inicio="3 días",
                        consejeria_alarma="Dificultad para respirar", consejeria_estado="Realizada",
                        recomendaciones="Reposo e hidratación", fecha_control="17/03/2026")
        raw["exam"] = {"general": {"estado": V.EXAM_NORMAL, "hallazgos": "Alerta, orientado"},
                       "abdomen": {"estado": V.EXAM_LIMITADO, "hallazgos": V.ABDOMEN_LIMITADO},
                       "torax": {"estado": V.EXAM_ANORMAL, "hallazgos": "Sibilancias espiratorias"}}
        raw["dx"] = [{"tipo": "principal", "descripcion": "Faringitis aguda", "cie11": "CA02.0", "cie11_verificado": 1},
                     {"tipo": "secundario", "descripcion": "Rinitis alérgica"}]
        raw["meds"] = [{"generico": "Acetaminofén", "presentacion": "Tab 500 mg", "dosis": "500 mg", "via": "Oral",
                        "frecuencia": "c/8 h", "duracion": "3 días", "cant_prescrita": "9", "cant_entregada": "9",
                        "unidad": "tabletas"}]
        data, errors, _ = V.prepare_visit("aps", raw)
        self.assertFalse(errors)
        aps_id = V.save_visit(self.db, "aps", ids[0], data)
        ssr_raw = raw_visit("ssr", "10/03/2026", personal="Dra. R", motivo="Planificación", categoria="Planificación familiar",
                            metodo_entregado="Preservativos", metodo_cantidad="10", metodo_unidad="unidades")
        ssr_raw["dx"] = [{"tipo": "principal", "descripcion": "Planificación familiar"}]
        ssr_raw["interventions"] = {V.SSR_INTERVENCIONES[2]: {"estado": "Realizada"}}
        ssr_raw["tests"] = [{"prueba": "Prueba de embarazo", "estado": "Realizada", "resultado": "Negativa"}]
        ssr_id = V.save_visit(self.db, "ssr", ids[0], V.prepare_visit("ssr", ssr_raw)[0])

        f1 = pdf.ficha_integral_pdf(self.db, ids[0], out / "integral.pdf")
        f2 = pdf.aps_pdf(self.db, aps_id, out / "aps.pdf")
        f3 = pdf.ssr_pdf(self.db, ssr_id, out / "ssr.pdf")
        # una APS totalmente vacía también debe generarse
        vacia = V.save_visit(self.db, "aps", ids[1], V.prepare_visit("aps", raw_visit("aps", "12/03/2026"))[0])
        f4 = pdf.aps_pdf(self.db, vacia, out / "aps_vacia.pdf")
        for f in (f1, f2, f3, f4):
            blob = Path(f).read_bytes()
            self.assertTrue(blob.startswith(b"%PDF"), f)
            self.assertGreater(len(blob), 2500, f)
        if shutil.which("pdftotext"):
            def text(p):
                return subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True,
                                      text=True).stdout
            t_aps, t_vacia = text(f2), text(f4)
            self.assertIn("FICHA DE ATENCIÓN PRIMARIA", t_aps)
            self.assertIn("Faringitis aguda", t_aps)
            self.assertIn("Exploración abdominal limitada", t_aps)
            self.assertIn("120", t_aps)
            self.assertIn("Pendiente / no registrado", t_vacia)
            self.assertIn("BORRADOR", t_vacia)
            self.assertNotIn("Sin hallazgos patológicos", t_vacia)
            self.assertNotIn("blando", t_vacia.lower())
            self.assertIn("No registrado", t_vacia)          # signos vitales
            self.assertIn("Página 1 de", t_vacia)
            t_ssr = text(f3)
            self.assertIn("Preservativos", t_ssr)
            self.assertIn("Prueba de embarazo", t_ssr)
            t_int = text(f1)
            self.assertIn("P%06d" % ids[0], t_int)
            self.assertIn("No constituye un formulario oficial", t_int)

    def test_12_excel(self):
        self._populate()
        rep = stats.run_report(self.db, "Pacientes atendidos por día")
        path = excel.export_report(Path(self.dir) / "rep.xlsx", rep, "Brigada Test")
        from openpyxl import load_workbook
        ws = load_workbook(path).active
        self.assertEqual(ws["A1"].value, "Pacientes atendidos por día")
        self.assertEqual([c.value for c in ws[6]][:2], ["Fecha", "Pacientes únicos"])
        values = [[c.value for c in row] for row in ws.iter_rows(min_row=7, max_row=7 + len(rep["rows"]) - 1)]
        self.assertEqual(values[-1][0], "TOTAL")
        self.assertEqual(values[-1][4], 5)
        pac = stats.run_report(self.db, "Listado administrativo de pacientes")
        path2 = excel.export_report(Path(self.dir) / "pac.xlsx", pac)
        ws2 = load_workbook(path2).active
        self.assertEqual(ws2.max_row >= 9, True)


class T06_Persistencia(unittest.TestCase):
    def test_13_cerrar_y_reabrir(self):
        d = tempfile.mkdtemp(dir=_TMP)
        path = Path(d) / "p.db"
        db = Database(path)
        pid = make_patient(db)
        vid = V.save_visit(db, "aps", pid, V.prepare_visit("aps", raw_visit("aps", "10/03/2026", personal="Dr. A"))[0])
        db.set_setting("encabezado_institucion", "Mi Brigada")
        db.close()
        db2 = Database(path)
        self.assertEqual(P.get_patient(db2, pid)["nombres"], "María José")
        self.assertEqual(V.get_visit(db2, "aps", vid)["main"]["personal"], "Dr. A")
        self.assertEqual(db2.get_setting("encabezado_institucion"), "Mi Brigada")
        self.assertEqual(db2.scalar("SELECT COUNT(*) FROM catalogs WHERE category='departamento'"), 18)
        db2.close()

    def test_14b_respaldo_segunda_ubicacion(self):
        d = tempfile.mkdtemp(dir=_TMP)
        db = Database(Path(d) / "p.db")
        make_patient(db)
        extra = Path(d) / "usb" / "MDm"
        db.set_setting("respaldo_extra_dir", str(extra))
        for _ in range(3):
            backup.auto_backup(db, prefix="periodica", keep=2)
            time.sleep(1.1)
        copias = sorted(extra.glob("periodica_*.db"))
        self.assertEqual(len(copias), 2)                          # rota también en la segunda ubicación
        self.assertEqual(backup.validate_backup(copias[-1])["pacientes"], 1)
        self.assertEqual(db.get_setting("respaldo_extra_error"), "")
        # ubicación inaccesible: la copia local se hace igual y queda el aviso
        bloqueo = Path(d) / "archivo.txt"
        bloqueo.write_text("x")
        db.set_setting("respaldo_extra_dir", str(bloqueo / "sub"))
        local = backup.auto_backup(db)
        self.assertTrue(local.exists())
        self.assertNotEqual(db.get_setting("respaldo_extra_error"), "")

    def test_14_respaldo_y_restauracion(self):
        d = tempfile.mkdtemp(dir=_TMP)
        db = Database(Path(d) / "p.db")
        pid = make_patient(db)
        V.save_visit(db, "aps", pid, V.prepare_visit("aps", raw_visit("aps", "10/03/2026"))[0])
        bk = backup.create_backup(db, Path(d) / "copia.db")
        self.assertTrue(bk.exists())
        info = backup.validate_backup(bk)
        self.assertEqual((info["pacientes"], info["aps"], info["ssr"]), (1, 1, 0))
        # se crean más datos después de la copia
        make_patient(db, "Pedro", "Soto")
        make_patient(db, "Rosa", "Mena")
        self.assertEqual(P.count_patients(db), 3)
        stats_restored = backup.restore_backup(db, bk)
        self.assertEqual(stats_restored["pacientes"], 1)
        self.assertEqual(P.count_patients(db), 1)
        self.assertEqual(P.get_patient(db, pid)["nombres"], "María José")
        # la base sigue funcionando tras restaurar
        make_patient(db, "Nuevo", "Paciente")
        self.assertEqual(P.count_patients(db), 2)
        # se creó la copia previa a la restauración
        self.assertTrue(any(p.name.startswith("previa_restauracion") for p in backup.list_backups()))
        # archivos inválidos se rechazan
        bad = Path(d) / "malo.db"
        bad.write_bytes(b"esto no es sqlite" * 50)
        with self.assertRaises(backup.BackupError):
            backup.validate_backup(bad)
        with self.assertRaises(backup.BackupError):
            backup.restore_backup(db, bad)
        self.assertEqual(P.count_patients(db), 2, "una restauración fallida no debe tocar los datos")
        empty = Path(d) / "otra.db"
        import sqlite3
        c = sqlite3.connect(empty); c.execute("CREATE TABLE x(a)"); c.commit(); c.close()
        with self.assertRaises(backup.BackupError):
            backup.validate_backup(empty)
        db.close()

    def test_catalogos_configurables(self):
        d = tempfile.mkdtemp(dir=_TMP)
        db = Database(Path(d) / "p.db")
        self.assertIn("Migrante retornado", db.catalog("estado_migratorio"))
        self.assertTrue(db.add_catalog("estado_migratorio", "Refugiado reasentado"))
        self.assertFalse(db.add_catalog("estado_migratorio", "Refugiado reasentado"))
        self.assertIn("Refugiado reasentado", db.catalog("estado_migratorio"))
        db.remove_catalog("estado_migratorio", "Refugiado reasentado")
        self.assertNotIn("Refugiado reasentado", db.catalog("estado_migratorio"))
        db.close()

    def test_contrasena(self):
        import security
        h = security.hash_password("clave-segura")
        self.assertNotIn("clave-segura", h)
        self.assertTrue(security.verify_password("clave-segura", h))
        self.assertFalse(security.verify_password("otra", h))
        self.assertFalse(security.verify_password("x", "basura"))


def tearDownModule():
    shutil.rmtree(_TMP, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
